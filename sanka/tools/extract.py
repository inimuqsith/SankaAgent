"""Text extraction shared by ``doc_read`` and the attachment tools.

The functions here never resolve user paths; callers validate the file first.
``read_*`` return ``(text, units, truncated)`` for one-shot reads.
``build_document`` extracts the whole document (with section offsets) and is
cached in memory so search, outline, and read share the exact same text.
Every extractor raises ``ToolError`` with a safe message on failure.
"""

from __future__ import annotations

import csv
import re
import threading
from collections import OrderedDict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sanka.tools.registry import ToolError

_MAX_DOC_PAGES = 300
_MAX_DOC_ROWS = 5_000
_MAX_DOC_SHEETS = 20
_MAX_DOC_CHARS = 2_000_000
_DOC_CACHE_SIZE = 4
_LINE_CHUNK = 200
_HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+)$", re.MULTILINE)


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return str(value)


def read_pdf(path: Path, max_pages: int) -> tuple[str, int, bool]:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ToolError("not_configured", "Pembaca PDF tidak tersedia.") from exc

    try:
        reader = PdfReader(str(path))
        total_pages = len(reader.pages)
        pages_to_read = min(max_pages, total_pages)
        parts: list[str] = []
        for page in reader.pages[:pages_to_read]:
            extracted = page.extract_text()
            parts.append(extracted if isinstance(extracted, str) else "")
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError("upstream_error", "Berkas PDF tidak dapat dibaca.") from exc

    text = "\n".join(parts).strip()
    return text, pages_to_read, total_pages > pages_to_read


def read_xlsx(path: Path, max_rows: int) -> tuple[str, int, bool]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ToolError("not_configured", "Pembaca Excel tidak tersedia.") from exc

    try:
        workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ToolError("upstream_error", "Berkas Excel tidak dapat dibaca.") from exc

    try:
        sheet = workbook.active
        if sheet is None:
            return "", 0, False
        lines: list[str] = []
        truncated = False
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            if index >= max_rows:
                truncated = True
                break
            lines.append("\t".join(_cell_text(value) for value in row))
        return "\n".join(lines), len(lines), truncated
    finally:
        workbook.close()


def read_csv(path: Path, max_rows: int) -> tuple[str, int, bool]:
    lines: list[str] = []
    truncated = False
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for index, row in enumerate(reader):
                if index >= max_rows:
                    truncated = True
                    break
                lines.append("\t".join(row))
    except OSError as exc:
        raise ToolError("upstream_error", "Berkas CSV tidak dapat dibaca.") from exc
    return "\n".join(lines), len(lines), truncated


def read_text(path: Path) -> tuple[str, int, bool]:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ToolError("upstream_error", "Berkas teks tidak dapat dibaca.") from exc
    return text, 1, False


def truncate(text: str, max_chars: int) -> tuple[str, bool]:
    if max_chars < 0:
        return text, False
    if len(text) <= max_chars:
        return text, False
    return text[:max_chars], True


@dataclass(frozen=True, slots=True)
class Section:
    """A labelled span of ``DocumentText.text``; offsets are string indices."""

    label: str
    start: int
    end: int


@dataclass(frozen=True, slots=True)
class DocumentText:
    """One full extraction shared by read, search, and outline."""

    text: str
    units: int
    truncated: bool
    sections: tuple[Section, ...]
    columns: tuple[str, ...] = ()


def _clean_label(value: str, limit: int = 80) -> str:
    cleaned = " ".join(value.split())
    return cleaned[:limit] or "bagian"


def _assemble(parts: list[str], labels: list[str]) -> tuple[str, tuple[Section, ...]]:
    """Join chunks with newlines and record each chunk's span in the result."""

    chunks: list[str] = []
    sections: list[Section] = []
    offset = 0
    for label, part in zip(labels, parts):
        if chunks:
            offset += 1
        start = offset
        chunks.append(part)
        offset += len(part)
        sections.append(Section(label=_clean_label(label), start=start, end=offset))
    return "\n".join(chunks), tuple(sections)


def _limit_document(
    text: str,
    sections: tuple[Section, ...],
    truncated: bool,
) -> tuple[str, tuple[Section, ...], bool]:
    """Cap the full text while keeping section offsets valid for the result."""

    if len(text) <= _MAX_DOC_CHARS:
        return text, sections, truncated
    cut = text[:_MAX_DOC_CHARS]
    kept = tuple(
        Section(label=section.label, start=section.start, end=min(section.end, len(cut)))
        for section in sections
        if section.start < len(cut)
    )
    return cut, kept, True


def _heading_sections(text: str) -> tuple[Section, ...] | None:
    """Return Markdown heading spans, or ``None`` when there is no heading."""

    matches = list(_HEADING_PATTERN.finditer(text))
    if not matches:
        return None
    sections: list[Section] = []
    if matches[0].start() > 0:
        sections.append(Section(label="pembuka", start=0, end=matches[0].start()))
    for index, match in enumerate(matches):
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        sections.append(
            Section(label=_clean_label(match.group(2)), start=match.start(), end=end)
        )
    return tuple(sections)


def _line_sections(text: str, chunk: int = _LINE_CHUNK) -> tuple[Section, ...]:
    """Split plain text into labelled line ranges of ``chunk`` lines."""

    sections: list[Section] = []
    start = 0
    line = 1
    length = len(text)
    while start < length:
        position = start
        count = 0
        while count < chunk:
            newline = text.find("\n", position)
            if newline == -1:
                position = length
                break
            position = newline + 1
            count += 1
        count = max(count, 1)
        sections.append(
            Section(label=f"baris {line}-{line + count - 1}", start=start, end=position)
        )
        line += count
        start = position
    return tuple(sections)


def _build_pdf(path: Path) -> DocumentText:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ToolError("not_configured", "Pembaca PDF tidak tersedia.") from exc

    try:
        reader = PdfReader(str(path))
        total_pages = len(reader.pages)
        pages_to_read = min(_MAX_DOC_PAGES, total_pages)
        parts: list[str] = []
        labels: list[str] = []
        for number, page in enumerate(reader.pages[:pages_to_read], start=1):
            extracted = page.extract_text()
            parts.append((extracted if isinstance(extracted, str) else "").strip())
            labels.append(f"halaman {number}")
    except ToolError:
        raise
    except Exception as exc:
        raise ToolError("upstream_error", "Berkas PDF tidak dapat dibaca.") from exc

    text, sections = _assemble(parts, labels)
    text, sections, truncated = _limit_document(
        text, sections, total_pages > pages_to_read
    )
    return DocumentText(
        text=text,
        units=pages_to_read,
        truncated=truncated,
        sections=sections,
    )


def _build_xlsx(path: Path) -> DocumentText:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency is declared
        raise ToolError("not_configured", "Pembaca Excel tidak tersedia.") from exc

    try:
        handle = path.open("rb")
    except OSError as exc:
        raise ToolError("upstream_error", "Berkas Excel tidak dapat dibaca.") from exc
    try:
        try:
            # Stored attachments have no extension, so openpyxl cannot infer the
            # format from the path; a binary stream makes it detect the ZIP.
            workbook = load_workbook(filename=handle, read_only=True, data_only=True)
        except Exception as exc:
            raise ToolError("upstream_error", "Berkas Excel tidak dapat dibaca.") from exc

        try:
            sheets = workbook.worksheets
            sheets_to_read = min(_MAX_DOC_SHEETS, len(sheets))
            truncated = len(sheets) > sheets_to_read
            parts: list[str] = []
            labels: list[str] = []
            columns: tuple[str, ...] = ()
            units = 0
            for sheet in sheets[:sheets_to_read]:
                lines: list[str] = []
                for index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if index >= _MAX_DOC_ROWS:
                        truncated = True
                        break
                    cells = [_cell_text(value) for value in row]
                    if not columns and any(cells):
                        columns = tuple(cells)
                    lines.append("\t".join(cells))
                units += len(lines)
                parts.append(f"### Lembar: {sheet.title}\n" + "\n".join(lines))
                labels.append(f"lembar {sheet.title}")
            text, sections = _assemble(parts, labels)
            text, sections, truncated = _limit_document(text, sections, truncated)
            return DocumentText(
                text=text,
                units=units,
                truncated=truncated,
                sections=sections,
                columns=columns,
            )
        finally:
            workbook.close()
    finally:
        handle.close()


def _build_csv(path: Path) -> DocumentText:
    lines: list[str] = []
    truncated = False
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.reader(handle)
            for index, row in enumerate(reader):
                if index >= _MAX_DOC_ROWS:
                    truncated = True
                    break
                lines.append("\t".join(row))
    except OSError as exc:
        raise ToolError("upstream_error", "Berkas CSV tidak dapat dibaca.") from exc

    columns = tuple(lines[0].split("\t")) if lines else ()
    parts: list[str] = []
    labels: list[str] = []
    for index in range(0, len(lines), _LINE_CHUNK):
        chunk = lines[index : index + _LINE_CHUNK]
        parts.append("\n".join(chunk))
        labels.append(f"baris {index + 1}-{index + len(chunk)}")
    text, sections = _assemble(parts, labels)
    text, sections, truncated = _limit_document(text, sections, truncated)
    return DocumentText(
        text=text,
        units=len(lines),
        truncated=truncated,
        sections=sections,
        columns=columns,
    )


def _build_text(path: Path) -> DocumentText:
    try:
        text = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ToolError("upstream_error", "Berkas teks tidak dapat dibaca.") from exc

    sections = _heading_sections(text)
    if sections is None:
        sections = _line_sections(text)
    text, sections, truncated = _limit_document(text, sections, False)
    units = text.count("\n") + 1 if text else 0
    return DocumentText(text=text, units=units, truncated=truncated, sections=sections)


_BUILDERS = {
    "pdf": _build_pdf,
    "xlsx": _build_xlsx,
    "csv": _build_csv,
    "text": _build_text,
}

_DOC_CACHE: OrderedDict[tuple[str, str, int, int], DocumentText] = OrderedDict()
_DOC_CACHE_LOCK = threading.Lock()


def build_document(path: Path, kind: str) -> DocumentText:
    """Extract ``path`` once and reuse the result for read/search/outline.

    The cache key includes size and ``mtime_ns`` so a replaced file is never
    served from a stale entry. Only four documents are kept in memory.
    """

    builder = _BUILDERS.get(kind)
    if builder is None:
        raise ToolError("invalid_arguments", "Jenis dokumen tidak didukung.")
    try:
        info = path.stat()
    except OSError as exc:
        raise ToolError("not_found", "Lampiran tidak ditemukan.") from exc

    key = (str(path), kind, info.st_size, info.st_mtime_ns)
    with _DOC_CACHE_LOCK:
        cached = _DOC_CACHE.get(key)
        if cached is not None:
            _DOC_CACHE.move_to_end(key)
            return cached

    document = builder(path)
    with _DOC_CACHE_LOCK:
        _DOC_CACHE[key] = document
        while len(_DOC_CACHE) > _DOC_CACHE_SIZE:
            _DOC_CACHE.popitem(last=False)
    return document
